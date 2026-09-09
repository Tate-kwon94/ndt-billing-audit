"""정답표(라벨링 시트) — 한 회차의 프로그램 결과를 뼈대로 층화 표본을 뽑아 정답 빈칸이 있는 엑셀을 만든다.

2026-09-07 사용자: "정답을 우리가 만들어놓고 사내에서 모자라는 부분을 메우는 방식". 흐름:
  1) ① 검토를 돌린다 (mock 이든 실 HCX 든) → billing_items·matches·inspection_reports·findings
  2) build_rows: 행마다 청구·매칭·성적서 쪽·판정·근거를 한 줄로 합치고 층 키를 단다
  3) sample: 층(판정×매칭방법×재확인×성적서유무)마다 고르게 n 행 (seed 고정 → 재현)
  4) write: 정답·근거·확신도·사내확인 빈칸 열이 있는 엑셀 + 층별 요약 시트
정답 열은 사람(또는 사외 검토)이 채우고, '사내확인_필요' 가 예인 행만 사내에서 본다.
"""
from __future__ import annotations

import random
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.database.models import BillingItem, Finding, InspectionReport, Match

PROGRAM_COLUMNS = ["행번호", "청구번호", "joint", "line", "NDT", "도면번호", "청구_성적서번호", "청구_결과", "수량",
                   "매칭_성적서번호", "매칭_방법", "매칭_점수", "성적서_쪽", "판정", "위험도", "재확인", "재확인_사유", "근거", "층"]
ANSWER_COLUMNS = ["정답_매칭성적서", "정답_판정", "정답_근거(문서/쪽)", "확신도", "사내확인_필요", "메모"]


def build_rows(session: Session, billing_round_id: int) -> list[dict]:
    items = (session.query(BillingItem).filter(BillingItem.billing_round_id == billing_round_id)
             .order_by(BillingItem.row_index).all())
    matches = {m.billing_item_id: m for m in session.query(Match)
               .filter(Match.billing_item_id.in_([i.id for i in items] or [-1])).all()}
    findings = {f.billing_item_id: f for f in session.query(Finding)
                .filter(Finding.billing_item_id.in_([i.id for i in items] or [-1])).all()}
    rep_ids = [m.inspection_report_id for m in matches.values() if m.inspection_report_id]
    reports = {r.id: r for r in session.query(InspectionReport)
               .filter(InspectionReport.id.in_(rep_ids or [-1])).all()}
    rows: list[dict] = []
    for it in items:
        m = matches.get(it.id)
        f = findings.get(it.id)
        rep = reports.get(m.inspection_report_id) if m and m.inspection_report_id else None
        review = bool((m and m.needs_review) or (f and f.needs_review))
        reasons = []
        for src in (m, f):
            if src is not None and src.review_reasons_json:
                reasons += list(src.review_reasons_json.get("reasons") or [])
        cites = (f.citations_json or {}).get("sources") if f else None
        evidence = "; ".join(f"{c.get('doc', '')} p.{c.get('page', '')}" for c in (cites or []) if isinstance(c, dict))
        verdict = f.verdict if f else "(판정 없음)"
        method = m.match_method if m else "none"
        rows.append({
            "행번호": it.row_index + 1, "청구번호": it.billing_no or "", "joint": it.joint_no, "line": it.line_no or "",
            "NDT": it.ndt_method, "도면번호": it.drawing_no or "", "청구_성적서번호": it.report_no or "",
            "청구_결과": it.result or "", "수량": it.quantity,
            "매칭_성적서번호": rep.report_no if rep else "", "매칭_방법": method,
            "매칭_점수": m.match_score if m else None,
            "성적서_쪽": f"p.{rep.start_page}-{rep.end_page}" if rep else "",
            "판정": verdict, "위험도": f.risk_score if f else None,
            "재확인": "예" if review else "", "재확인_사유": " / ".join(reasons), "근거": evidence,
            "층": "|".join([verdict, method, "재확인" if review else "-", "성적서" if rep else "성적서없음"]),
        })
    return rows


def sample(rows: list[dict], *, n: int, seed: int = 0) -> list[dict]:
    """층마다 최소 1행, 나머지는 층 크기에 비례해 배분. 같은 seed 면 같은 결과."""
    by = defaultdict(list)
    for r in rows:
        by[r["층"]].append(r)
    strata = sorted(by)
    if not strata:
        return []
    rng = random.Random(seed)
    quota = {k: 1 for k in strata}
    remaining = max(0, n - len(strata))
    total = sum(len(v) for v in by.values())
    for k in strata:
        quota[k] += int(remaining * len(by[k]) / total) if total else 0
    # 반올림으로 남은 자리는 큰 층부터
    while sum(min(quota[k], len(by[k])) for k in strata) < min(n, total):
        for k in sorted(strata, key=lambda k: -len(by[k])):
            if quota[k] < len(by[k]):
                quota[k] += 1
                break
        else:
            break
    picked: list[dict] = []
    for k in strata:
        pool = sorted(by[k], key=lambda r: r["행번호"])
        picked += rng.sample(pool, min(quota[k], len(pool)))
    picked.sort(key=lambda r: r["행번호"])
    return picked[:n]


def write(rows: list[dict], path: Path, *, round_label: str = "") -> Path:
    path = Path(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "정답표"
    header = PROGRAM_COLUMNS + ANSWER_COLUMNS
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for j, name in enumerate(ANSWER_COLUMNS, start=len(PROGRAM_COLUMNS) + 1):
        ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor="FFF2CC")
    for r in rows:
        ws.append([r.get(k) for k in PROGRAM_COLUMNS] + [""] * len(ANSWER_COLUMNS))
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJKLMNOPQRS", (6, 10, 14, 16, 5, 30, 12, 6, 6, 12, 8, 7, 10, 13, 6, 6, 40, 30, 28)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("층별_요약")
    ws2.append(["회차", round_label]); ws2.append([])
    ws2.append(["층(판정|매칭|재확인|성적서)", "표본 행수"])
    for k, v in sorted(Counter(r["층"] for r in rows).items()):
        ws2.append([k, v])
    ws2.append([]); ws2.append(["채우는 법"])
    ws2.append(["정답_판정", "OK / SUSPECT / NONCOMPLIANT 중 하나. 프로그램 판정과 같으면 그대로 적는다."])
    ws2.append(["정답_근거(문서/쪽)", "판단의 근거가 된 문서와 쪽. 예: 성적서 12-056 p.41, SP 70 p.88"])
    ws2.append(["확신도", "높음 / 중간 / 낮음. 낮음이면 사내확인_필요 = 예"])
    ws2.append(["사내확인_필요", "예 = 사외에서 못 읽은 것(손글씨·도장 위 글자·원본 필요). 사내에서 이 행만 본다."])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
