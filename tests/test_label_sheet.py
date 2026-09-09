"""정답표(라벨링 시트) — 1항차 실 데이터로 사외에서 정답을 만들고 사내에서 빈칸을 메우는 방식 (2026-09-07 사용자).

프로그램이 뽑은 매칭·판정·근거 쪽수를 뼈대로, 층(판정×매칭방법×재확인×성적서유무)마다 고르게 표본을 뽑아
정답·근거·확신도·사내확인 빈칸이 있는 엑셀을 만든다. 이 표가 이후 모든 A/B 의 기준선.
"""
from datetime import date

import openpyxl

from app.report import label_sheet as ls


def _seed(models, repo, s):
    br = repo.create_billing_round(s, round_no=1, discipline="CP-P1", billing_date=date(2026, 8, 30),
                                   reports_pdf_path="C:/x/reports.pdf")
    rows = []
    for i in range(40):
        rows.append({"joint_no": f"J{i:03d}", "line_no": "20PAB02BR005", "ndt_method": "VT" if i % 2 else "PT",
                     "report_no": f"12-{i:03d}", "result": "A", "quantity": 1.0})
    repo.add_billing_items(s, br.id, rows)
    s.flush()
    items = s.query(models.BillingItem).filter_by(billing_round_id=br.id).order_by(models.BillingItem.row_index).all()
    for i, it in enumerate(items):
        rep = None
        if i % 4 != 3:                                     # 4행 중 3행은 성적서 있음
            rep = repo.add_inspection_report(s, billing_round_id=br.id, source_pdf="C:/x/reports.pdf",
                                             start_page=10 + i, end_page=11 + i, report_no=f"12-{i:03d}", ndt_method="VT")
            s.flush()
        repo.save_match(s, billing_item_id=it.id, inspection_report_id=rep.id if rep else None,
                        matched_joint_no=it.joint_no if rep else None,
                        match_method=("exact" if i % 3 == 0 else "fuzzy") if rep else "none",
                        match_score=0.9 if rep else None, reasoning="r", needs_review=(i % 5 == 0),
                        review_reasons_json={"reasons": ["후보 점수 근접"]} if i % 5 == 0 else None)
        repo.save_finding(s, billing_item_id=it.id, verdict=["OK", "SUSPECT", "NONCOMPLIANT"][i % 3], risk_score=i % 10,
                          summary="s", citations_json={"sources": [{"doc": "SP 70", "page": 88}]},
                          recommended_action="a", needs_review=(i % 5 == 0))
    s.commit()
    return br.id


def test_rows_join_item_match_report_finding_with_stratum(fresh_db):
    models, repo = fresh_db
    with models.get_session() as s:
        rid = _seed(models, repo, s)
        rows = ls.build_rows(s, rid)
    assert len(rows) == 40
    r0 = rows[0]
    for k in ("행번호", "joint", "NDT", "청구_성적서번호", "매칭_성적서번호", "매칭_방법", "성적서_쪽", "판정", "위험도",
              "재확인", "재확인_사유", "근거", "층"):
        assert k in r0, k
    assert r0["성적서_쪽"] == "p.10-11"
    assert rows[3]["매칭_방법"] == "none" and rows[3]["성적서_쪽"] == ""
    assert r0["층"] == "OK|exact|재확인|성적서"


def test_stratified_sample_covers_every_stratum_and_caps_total(fresh_db):
    models, repo = fresh_db
    with models.get_session() as s:
        rid = _seed(models, repo, s)
        rows = ls.build_rows(s, rid)
    picked = ls.sample(rows, n=15, seed=1)
    assert len(picked) <= 15
    assert {r["층"] for r in picked} == {r["층"] for r in rows}       # 모든 층이 최소 1행
    assert ls.sample(rows, n=15, seed=1) == picked                       # 같은 seed → 같은 표본 (재현)


def test_write_sheet_has_answer_columns_and_summary(fresh_db, tmp_path):
    models, repo = fresh_db
    with models.get_session() as s:
        rid = _seed(models, repo, s)
        rows = ls.build_rows(s, rid)
    out = tmp_path / "labels.xlsx"
    ls.write(ls.sample(rows, n=20, seed=1), out, round_label="1항차 CP-P1")
    wb = openpyxl.load_workbook(out)
    assert set(wb.sheetnames) >= {"정답표", "층별_요약"}
    ws = wb["정답표"]
    header = [c.value for c in ws[1]]
    for k in ("정답_매칭성적서", "정답_판정", "정답_근거(문서/쪽)", "확신도", "사내확인_필요", "메모"):
        assert k in header, k
    assert ws.max_row >= 2


def test_cli_label_sheet(fresh_db, tmp_path):
    from typer.testing import CliRunner
    from app.main import app
    models, repo = fresh_db
    with models.get_session() as s:
        _seed(models, repo, s)
    out = tmp_path / "l.xlsx"
    r = CliRunner().invoke(app, ["label-sheet", "--round", "1", "--n", "12", "--out", str(out)])
    assert r.exit_code == 0, r.output
    assert out.exists() and "12" in r.output
