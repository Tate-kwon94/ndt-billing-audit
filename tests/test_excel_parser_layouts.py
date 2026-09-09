"""청구 엑셀 양식 회귀 — 파서 단위 테스트 (2026-09-06 이전에는 0건).

2026-09-06 발견: 시공사가 지금 내는 첨부-2 양식(0827·0831)을 파서가 못 읽었다.
  - 헤더 판정 임계가 '후보 별칭 수 // 4' 라 별칭을 늘릴수록 못 찾는 구조
  - 용접부 열이 'Item' 인데 joint_no 후보에 없음
  - 도면 연결이 'Weld Map No.' 인데 시트별 보강이 'Detailed Drawing'/'Welding Map' 만 찾음
옛 양식(BLDG·Confirmation No.)은 계속 읽혀야 한다.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from app.extractors.excel_parser import parse_billing_xlsx, enrich_from_per_method_sheets


def _new_layout(path: Path) -> Path:
    """첨부-2 양식: 3행 헤더(줄바꿈 포함), 'CP-P1(전체)' + 'VT(UNIT1)' 시트, Weld Map No. 는 X열."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CP-P1(전체)"
    hdr = ["No", "Unit", "Application No.", "Report Number", "Date of Testing", "Type", "Item", "Result",
           "Dimension\n(Ø×t, mm)", "Thickness\n(mm)", "VT"] + [None] * 6 + ["PT Point "] + [None] * 2 + ["UT Point"] + [None] * 2 + ["Weld Map No."]
    ws.append(hdr)
    ws.append(["PIPE", "Rebar & etc", "Total\nQ'ty"])
    ws.append(["T≤4", ' 4"＜ t ≤10"'])
    wm = "NP.008.CCW.ABD.1.021.0004"
    ws.append([1, "1", "NDT-00010", "77-005PT", datetime(2024, 2, 27), "PT", "FW12", "A", "3040x20", 20] + [None] * 13 + [wm])
    ws.append([2, "1", "NDT-00010", "77-013VMC", datetime(2024, 2, 27), "VT", "FW12", "A", "3040x20", 20] + [None] * 13 + [wm])
    ws.append([3, "2", "NDT-00102", "77-083VMC", datetime(2024, 5, 12), "VT", "(1÷65)SC2-1", "A", "Support", 5] + [None] * 13 + [None])
    vt = wb.create_sheet("VT(UNIT1)")
    vt.append(["No", "Unit", "Application \nNo.", "Report \nNumber", "Date of \nTesting", "Type", "Item", "Result",
               "Dimension\n(Ø×t, mm", "Thick\n(mm)"] + [None] * 7 + ["Weld Map No."])
    vt.append(["PIPE", "Rebar & etc", "Total\nQ'ty"])
    vt.append(["T≤4"])
    vt.append([2, "1", "NDT-00010", "77-013VMC", "27.02.2024", "VT", "FW12", "A", "3040x20", 20] + [None] * 7 + [wm])
    wb.save(path)
    return path


def _old_layout(path: Path) -> Path:
    """2024-12 판 양식: BLDG · Welder ID · Confirmation No."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CP-P1(Total)"
    ws.append(["◇ CP-P1(Total)"])
    ws.append(["No", "Unit", "BLDG", "Application No", "Report Number", "Date of Testing", "Type", "Result",
               "Welder ID", "Confirmation No."])
    ws.append(["Pipe", "Rebar & etc", "Total\n(Q'ty)"])
    ws.append([1, "1", "10UMA", "NDT-00010", "77-005PT", datetime(2024, 2, 19), "PT", "A", "DW003 DW005", "FW12"])
    wb.save(path)
    return path


def test_new_layout_header_is_detected_and_item_is_joint(tmp_path):
    p = _new_layout(tmp_path / "new.xlsx")
    parsed = parse_billing_xlsx(p, discipline_hint="CP-P1")
    assert parsed.sheet_name == "CP-P1(전체)"
    assert parsed.header_row == 1
    assert len(parsed.rows) == 3
    assert [r["joint_no"] for r in parsed.rows] == ["FW12", "FW12", "(1÷65)SC2-1"]
    assert [r["ndt_method"] for r in parsed.rows] == ["PT", "VT", "VT"]
    assert parsed.rows[0]["report_no"] == "77-005PT"
    # 2~3행 부헤더('PIPE…', 'T≤4…')는 joint/method 가 없어 _invalid 로 걸러진다 — 기존 동작, 조용히 사라지지 않고 센다
    assert len(parsed.skipped_invalid_rows) == 2


def test_old_layout_still_parses(tmp_path):
    p = _old_layout(tmp_path / "old.xlsx")
    parsed = parse_billing_xlsx(p, discipline_hint="CP-P1")
    assert len(parsed.rows) == 1
    assert parsed.rows[0]["joint_no"] == "FW12"
    assert parsed.rows[0]["bldg"] == "10UMA"


def test_new_layout_weld_map_is_enriched(tmp_path):
    p = _new_layout(tmp_path / "new.xlsx")
    en = enrich_from_per_method_sheets(p, discipline="CP-P1")
    assert "VT(UNIT1)" in en.source_sheets
    assert en.by_report_no["77-013VMC"]["welding_map"] == "NP.008.CCW.ABD.1.021.0004"
    assert en.by_report_no["77-013VMC"]["drawing_no"] is None


def test_adding_aliases_does_not_break_header_detection(tmp_path, monkeypatch):
    """임계가 별칭 수에 비례하면 별칭을 늘리는 순간 헤더를 못 찾는다. 필드 수 기준이어야 한다."""
    from app.config import templates_config
    cfg = templates_config()
    spec = cfg["disciplines"]["CP-P1"]["columns"]
    for col in spec:
        col["candidates"] = list(col["candidates"]) + [f"alias{i}" for i in range(20)]
    try:
        p = _new_layout(tmp_path / "new.xlsx")
        parsed = parse_billing_xlsx(p, discipline_hint="CP-P1")
        assert len(parsed.rows) == 3
    finally:
        for col in spec:
            col["candidates"] = [c for c in col["candidates"] if not str(c).startswith("alias")]


REAL = Path("samples/billing/20260830/첨부-2 the plant NDT Inpsection Service Detail List_Piping_20260831.xlsx")


@pytest.mark.skipif(not REAL.exists(), reason="실 청구 엑셀 없음 (사외 샘플 전용)")
def test_real_0831_billing_parses_with_joints():
    parsed = parse_billing_xlsx(REAL, discipline_hint="CP-P1")
    assert len(parsed.rows) > 4000
    with_joint = sum(1 for r in parsed.rows if r["joint_no"])
    assert with_joint / len(parsed.rows) > 0.95
    en = enrich_from_per_method_sheets(REAL, discipline="CP-P1")
    assert sum(1 for v in en.by_report_no.values() if v.get("welding_map")) > 1000
