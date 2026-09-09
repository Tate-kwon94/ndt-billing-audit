"""검토 엑셀 쓰기 end-to-end — 파싱한 청구 엑셀을 회차로 적재한 뒤 write() 가 실제로 파일을 낸다.

2026-09-07 스모크에서 발견: review 가 4,395행 검토를 다 마친 뒤 마지막 엑셀 쓰기에서 죽었다.
_find_header_row_and_cells 시그니처를 바꾸며 excel_writer 의 호출을 안 고쳤고, 기존 테스트는
행 조립 함수만 검사해 write() 의 헤더 재탐색 경로를 안 탔다. 이 파일이 그 경로를 고정한다.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl

from tests.test_excel_parser_layouts import _new_layout, _old_layout


def _round_from(m, r, xlsx: Path, *, round_no: int) -> int:
    from app.extractors.excel_parser import parse_billing_xlsx
    parsed = parse_billing_xlsx(xlsx, discipline_hint="CP-P1")
    with m.get_session() as s:
        br = r.create_billing_round(s, round_no=round_no, discipline="CP-P1", billing_date=date(2026, 8, 30),
                                    billing_xlsx_path=str(xlsx), reports_pdf_path="r.pdf")
        s.flush()
        assert r.add_billing_items(s, br.id, parsed.rows) == len(parsed.rows)
        s.commit()
        return br.id


def test_write_review_excel_for_new_layout(fresh_db, tmp_path):
    m, r = fresh_db
    xlsx = _new_layout(tmp_path / "new.xlsx")
    rid = _round_from(m, r, xlsx, round_no=25)
    from app.report.excel_writer import write
    out = Path(write(rid))
    assert out.exists() and out.suffix == ".xlsx"
    wb = openpyxl.load_workbook(out, read_only=True)
    ws = wb[wb.sheetnames[0]] if "CP-P1(전체)" not in wb.sheetnames else wb["CP-P1(전체)"]
    header = [str(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)) if c is not None]
    assert "적합성_판정" in header and "근거_상태" in header      # 검토 컬럼이 원본 헤더 행에 붙는다


def test_write_review_excel_for_old_layout(fresh_db, tmp_path):
    m, r = fresh_db
    xlsx = _old_layout(tmp_path / "old.xlsx")
    rid = _round_from(m, r, xlsx, round_no=1)
    from app.report.excel_writer import write
    out = Path(write(rid))
    assert out.exists()
